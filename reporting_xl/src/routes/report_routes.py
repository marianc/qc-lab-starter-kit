from fastapi import APIRouter, HTTPException
from fastapi.responses import Response
import traceback
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from ..dtos.report_dtos import TestingReportsExportRequestDto

router = APIRouter()

def register_report_routes(app):
    @app.post("/testing-reports-excel")
    def generate_testing_reports_excel(payload: TestingReportsExportRequestDto):
        try:
            wb = Workbook()
            # Remove default active sheet later once we create our sheets
            default_sheet = wb.active

            used_sheet_names = set()
            created_sheets = []

            for sheet_data in payload.sheets:
                # 1. Determine sheet label and ensure unique & reasonable length (max 31 chars for Excel)
                base_label = sheet_data.sheetLabel.strip()
                if len(base_label) > 31:
                    base_label = base_label[:31]
                
                sheet_name = base_label
                counter = 1
                while sheet_name.lower() in used_sheet_names:
                    suffix = f"_{counter}"
                    max_base_len = 31 - len(suffix)
                    sheet_name = base_label[:max_base_len] + suffix
                    counter += 1
                
                used_sheet_names.add(sheet_name.lower())

                ws = wb.create_sheet(title=sheet_name)
                ws.views.sheetView[0].showGridLines = True
                created_sheets.append((sheet_data, ws))

            # Remove default sheet if it exists and we created at least one sheet
            if len(wb.worksheets) > 1 and default_sheet in wb.worksheets:
                wb.remove(default_sheet)

            thin_border = Border(
                left=Side(style='thin', color='D9D9D9'),
                right=Side(style='thin', color='D9D9D9'),
                top=Side(style='thin', color='D9D9D9'),
                bottom=Side(style='thin', color='D9D9D9')
            )
            data_font = Font(name="Calibri", size=11)

            for sheet_data, ws in created_sheets:
                # 2. Title row
                ws.append([sheet_data.sheetTitle])
                total_cols = 4 + len(sheet_data.tests)
                ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
                title_cell = ws.cell(row=1, column=1)
                title_cell.font = Font(name="Calibri", size=14, bold=True)
                title_cell.alignment = Alignment(horizontal="left", vertical="center")
                ws.row_dimensions[1].height = 30

                # Blank row
                ws.append([])
                ws.row_dimensions[2].height = 15

                # 3. Header row
                third_col_header = "Material Name" if sheet_data.isCategory else "Control Code"
                headers = ["ID", "Submitted date", third_col_header, "Measurement ID"] + [t.name for t in sheet_data.tests]
                ws.append(headers)
                header_row_idx = 3
                ws.row_dimensions[header_row_idx].height = 35

                header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
                header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")

                for col_idx in range(1, len(headers) + 1):
                    cell = ws.cell(row=header_row_idx, column=col_idx)
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    cell.border = thin_border

                # 4. Data rows
                current_row_idx = 4

                # Group rows by reportId
                report_groups = []
                for row in sheet_data.rows:
                    if not report_groups or report_groups[-1]["reportId"] != row.reportId:
                        report_groups.append({
                            "reportId": row.reportId,
                            "dateSubmitted": row.dateSubmitted,
                            "controlCode": row.controlCode,
                            "materialName": row.materialName,
                            "measurements": []
                        })
                    report_groups[-1]["measurements"].append(row)

                for group in report_groups:
                    rep_id = group["reportId"]
                    
                    # Format date to dd.mm.yyyy if possible
                    date_sub_raw = group["dateSubmitted"]
                    date_sub_formatted = date_sub_raw
                    if date_sub_raw:
                        try:
                            # Try parsing yyyy-MM-dd HH:mm or similar
                            dt = datetime.fromisoformat(date_sub_raw.replace(' ', 'T'))
                            date_sub_formatted = dt.strftime("%d.%m.%Y")
                        except Exception:
                            try:
                                dt = datetime.strptime(date_sub_raw[:10], "%Y-%m-%d")
                                date_sub_formatted = dt.strftime("%d.%m.%Y")
                            except Exception:
                                pass

                    third_col_val = group["materialName"] if sheet_data.isCategory else group["controlCode"]

                    group_start_row = current_row_idx
                    group_total_physical_rows = 0

                    measurement_item_blocks = []
                    for meas in group["measurements"]:
                        max_lines = 1
                        test_val_map = {tv.testId: tv.values for tv in meas.testValues}
                        for t in sheet_data.tests:
                            vals = test_val_map.get(t.testId, [""])
                            if not vals:
                                vals = [""]
                            if len(vals) > max_lines:
                                max_lines = len(vals)
                        measurement_item_blocks.append({
                            "meas": meas,
                            "max_lines": max_lines,
                            "test_val_map": test_val_map
                        })
                        group_total_physical_rows += max_lines

                    group_end_row = group_start_row + group_total_physical_rows - 1

                    curr_meas_row = group_start_row
                    for block in measurement_item_blocks:
                        meas = block["meas"]
                        max_lines = block["max_lines"]
                        test_val_map = block["test_val_map"]
                        meas_id_str = f"{meas.measurementId}{'*' if meas.hasForm else ''}"

                        for line_idx in range(max_lines):
                            row_data = [
                                rep_id if line_idx == 0 and curr_meas_row == group_start_row else "",
                                date_sub_formatted if line_idx == 0 and curr_meas_row == group_start_row else "",
                                third_col_val if line_idx == 0 and curr_meas_row == group_start_row else "",
                                meas_id_str if line_idx == 0 else ""
                            ]

                            for t in sheet_data.tests:
                                vals = test_val_map.get(t.testId, [""])
                                if line_idx < len(vals):
                                    row_data.append(vals[line_idx])
                                else:
                                    row_data.append("")

                            ws.append(row_data)
                            ws.row_dimensions[curr_meas_row].height = 20

                            # Apply styling, borders and center alignment for all data cells
                            for c_idx in range(1, len(headers) + 1):
                                c = ws.cell(row=curr_meas_row, column=c_idx)
                                c.font = data_font
                                c.border = thin_border
                                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

                            curr_meas_row += 1

                    # Merge ID, DateSubmitted, ThirdCol vertically across the entire report group if group_total_physical_rows > 1
                    if group_total_physical_rows > 1:
                        ws.merge_cells(start_row=group_start_row, start_column=1, end_row=group_end_row, end_column=1)
                        ws.merge_cells(start_row=group_start_row, start_column=2, end_row=group_end_row, end_column=2)
                        ws.merge_cells(start_row=group_start_row, start_column=3, end_row=group_end_row, end_column=3)

                    # Merge Measurement ID (col 4) for each measurement block if it has multiple lines (array test)
                    meas_block_start = group_start_row
                    for block in measurement_item_blocks:
                        m_lines = block["max_lines"]
                        if m_lines > 1:
                            ws.merge_cells(start_row=meas_block_start, start_column=4, end_row=meas_block_start + m_lines - 1, end_column=4)
                        meas_block_start += m_lines

                    current_row_idx = curr_meas_row

                # 5. Column widths formatting
                ws.column_dimensions['A'].width = 12 # ID
                ws.column_dimensions['B'].width = 18 # Date Submitted
                ws.column_dimensions['C'].width = 25 # Control Code / Material Name
                ws.column_dimensions['D'].width = 18 # Measurement ID

                for c_idx in range(5, len(headers) + 1):
                    col_letter = get_column_letter(c_idx)
                    ws.column_dimensions[col_letter].width = 20

            # Save workbook to buffer
            from io import BytesIO
            bio = BytesIO()
            wb.save(bio)
            bio.seek(0)

            return Response(
                content=bio.getvalue(),
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={
                    "Content-Disposition": "attachment; filename=testing_reports.xlsx"
                }
            )
        except Exception as e:
            err_msg = "".join(traceback.format_exception(type(e), e, e.__traceback__))
            print("ERROR generating testing reports Excel:", err_msg)
            raise HTTPException(status_code=500, detail=str(e) + "\n" + err_msg)
